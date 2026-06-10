import csv
import io
import re
import time
import unicodedata
import uuid
from decimal import Decimal, InvalidOperation

from django.db import IntegrityError, transaction

from apps.products.models import Category, Product


MAX_IMPORT_ROWS = 5000
MAX_UPLOAD_BYTES = 5 * 1024 * 1024
SUPPORTED_EXTENSIONS = {".csv", ".xlsx", ".xls"}
REQUIRED_COLUMNS = {"name", "price"}
PRODUCT_DECIMAL_LIMITS = {
    "price": (12, 2),
    "cost": (12, 2),
    "stock": (12, 3),
    "min_stock": (12, 3),
}


def _clean_text(value, max_length=None):
    if value is None:
        text = ""
    elif isinstance(value, float) and value.is_integer():
        text = str(int(value))
    else:
        text = str(value)
    text = text.replace("\ufeff", "").strip()
    if max_length is not None:
        return text[:max_length]
    return text


def _normalize_key(value):
    text = _clean_text(value).lower()
    text = unicodedata.normalize("NFKD", text)
    text = "".join(char for char in text if not unicodedata.combining(char))
    text = re.sub(r"[^a-z0-9]+", "_", text)
    return re.sub(r"_+", "_", text).strip("_")


_COLUMN_ALIASES = {
    "name": "name",
    "nombre": "name",
    "producto": "name",
    "product": "name",
    "description": "description",
    "descripcion": "description",
    "detalle": "description",
    "price": "price",
    "precio": "price",
    "precio_venta": "price",
    "cost": "cost",
    "costo": "cost",
    "precio_costo": "cost",
    "unit": "unit",
    "unidad": "unit",
    "stock": "stock",
    "existencias": "stock",
    "quantity": "stock",
    "cantidad": "stock",
    "min_stock": "min_stock",
    "stock_minimo": "min_stock",
    "minimum_stock": "min_stock",
    "barcode": "barcode",
    "codigo": "barcode",
    "codigo_barras": "barcode",
    "codigo_de_barras": "barcode",
    "ean": "barcode",
    "sku": "barcode",
    "image_url": "image_url",
    "imagen": "image_url",
    "url_imagen": "image_url",
    "image": "image_url",
    "category": "category",
    "categoria": "category",
    "category_name": "category",
    "nombre_categoria": "category",
    "category_id": "category_id",
    "categoria_id": "category_id",
}

COLUMN_ALIASES = {_normalize_key(key): value for key, value in _COLUMN_ALIASES.items()}


class ProductImportFormatError(Exception):
    def __init__(self, message, code="invalid_file", errors=None):
        super().__init__(message)
        self.message = message
        self.code = code
        self.errors = errors or [{"row": None, "field": None, "code": code, "message": message}]


def build_failed_import_payload(message, code="invalid_file", errors=None, elapsed_ms=0):
    return {
        "result": "FAILED",
        "summary": {
            "total_rows": 0,
            "created": 0,
            "updated": 0,
            "failed": 0,
            "matches": 0,
            "elapsed_ms": elapsed_ms,
            "eta_ms": 0,
        },
        "items_loaded": [],
        "matches": [],
        "errors": errors or [{"row": None, "field": None, "code": code, "message": message}],
    }


def _row_error(row, field, message, code="invalid"):
    return {
        "row": row,
        "field": field,
        "code": code,
        "message": message,
    }


class ProductImportService:
    def preview_upload(self, tenant_id, upload):
        started = time.perf_counter()
        rows = self._read_upload(upload)
        if hasattr(upload, "seek"):
            upload.seek(0)

        created = 0
        updated = 0
        errors = []
        seen_file_keys = {}

        for row in rows:
            row_number = row["__row__"]
            product_data, row_errors = self._normalize_product_data(tenant_id, row)

            if not row_errors:
                row_errors.extend(
                    self._validate_file_duplicates(
                        row_number,
                        product_data,
                        seen_file_keys,
                    )
                )

            if not row_errors:
                row_errors.extend(
                    self._preview_category_errors(
                        tenant_id,
                        product_data.get("_category_ref"),
                        row_number,
                    )
                )

            match = None
            if not row_errors:
                match = self._match_active_product(tenant_id, product_data, row_number)
                row_errors.extend(match["errors"])

            if row_errors:
                errors.extend(row_errors)
                continue

            if match["product"] is None:
                created += 1
            else:
                updated += 1
            self._remember_file_keys(row_number, product_data, seen_file_keys)

        elapsed_ms = int((time.perf_counter() - started) * 1000)
        failed_rows = len({error["row"] for error in errors if error.get("row") is not None})
        return {
            "total_rows": len(rows),
            "created": created,
            "updated": updated,
            "failed": failed_rows,
            "errors": errors,
            "elapsed_ms": elapsed_ms,
        }

    def import_upload(self, tenant_id, upload):
        started = time.perf_counter()
        rows = self._read_upload(upload)
        loaded_items = []
        matches = []
        errors = []
        seen_file_keys = {}

        for row in rows:
            row_number = row["__row__"]
            product_data, row_errors = self._normalize_product_data(tenant_id, row)

            if not row_errors:
                duplicate_errors = self._validate_file_duplicates(
                    row_number,
                    product_data,
                    seen_file_keys,
                )
                row_errors.extend(duplicate_errors)

            if row_errors:
                errors.extend(row_errors)
                continue

            upsert_result = self._upsert_product(tenant_id, product_data, row_number)
            if upsert_result["errors"]:
                errors.extend(upsert_result["errors"])
                continue

            product = upsert_result["product"]
            loaded_items.append({
                "row": row_number,
                "action": upsert_result["action"],
                "match_type": upsert_result["match_type"],
                "product": product,
            })
            self._remember_file_keys(row_number, product_data, seen_file_keys)

            if upsert_result["match_type"]:
                matches.append({
                    "row": row_number,
                    "match_type": upsert_result["match_type"],
                    "action": upsert_result["action"],
                    "product_uuid": str(product.uuid),
                    "name": product.name,
                    "barcode": product.barcode,
                })

        elapsed_ms = int((time.perf_counter() - started) * 1000)
        failed_rows = len({error["row"] for error in errors if error.get("row") is not None})
        created = sum(1 for item in loaded_items if item["action"] == "created")
        updated = sum(1 for item in loaded_items if item["action"] == "updated")

        if loaded_items and not errors:
            result = "OK"
        elif loaded_items:
            result = "PARTIAL"
        else:
            result = "FAILED"

        return {
            "result": result,
            "summary": {
                "total_rows": len(rows),
                "created": created,
                "updated": updated,
                "failed": failed_rows,
                "matches": len(matches),
                "elapsed_ms": elapsed_ms,
                "eta_ms": 0,
            },
            "items": loaded_items,
            "matches": matches,
            "errors": errors,
        }

    def _read_upload(self, upload):
        filename = _clean_text(getattr(upload, "name", ""))
        extension = self._extension(filename)
        if extension not in SUPPORTED_EXTENSIONS:
            allowed = ", ".join(sorted(SUPPORTED_EXTENSIONS))
            raise ProductImportFormatError(
                f"Formato de archivo no soportado. Formatos permitidos: {allowed}.",
                code="unsupported_file_type",
            )

        upload_size = getattr(upload, "size", None)
        if upload_size is not None and upload_size > MAX_UPLOAD_BYTES:
            max_mb = MAX_UPLOAD_BYTES // (1024 * 1024)
            raise ProductImportFormatError(
                f"El archivo supera el maximo permitido de {max_mb} MB.",
                code="file_too_large",
            )

        if hasattr(upload, "seek"):
            upload.seek(0)

        if extension == ".csv":
            table_rows = self._read_csv(upload)
        elif extension == ".xlsx":
            table_rows = self._read_xlsx(upload)
        else:
            table_rows = self._read_xls(upload)

        return self._table_rows_to_dicts(table_rows)

    def _extension(self, filename):
        if "." not in filename:
            return ""
        return "." + filename.rsplit(".", 1)[1].lower()

    def _read_csv(self, upload):
        raw = upload.read()
        if not raw:
            raise ProductImportFormatError("El archivo esta vacio.", code="empty_file")
        text = self._decode_csv(raw)
        if not text.strip():
            raise ProductImportFormatError("El archivo esta vacio.", code="empty_file")
        try:
            dialect = csv.Sniffer().sniff(text[:2048])
        except csv.Error:
            dialect = csv.excel
        rows = []
        for row in csv.reader(io.StringIO(text), dialect):
            rows.append(row)
            if len(rows) > MAX_IMPORT_ROWS + 1:
                break
        return rows

    def _decode_csv(self, raw):
        for encoding in ("utf-8-sig", "latin-1"):
            try:
                return raw.decode(encoding)
            except UnicodeDecodeError:
                continue
        raise ProductImportFormatError(
            "No se pudo decodificar el CSV. Usar UTF-8 o Latin-1.",
            code="invalid_encoding",
        )

    def _read_xlsx(self, upload):
        try:
            from openpyxl import load_workbook
        except ModuleNotFoundError as exc:
            raise ProductImportFormatError(
                "El soporte para .xlsx requiere instalar openpyxl.",
                code="missing_dependency",
            ) from exc

        try:
            workbook = load_workbook(upload, read_only=True, data_only=True)
        except Exception as exc:
            raise ProductImportFormatError(
                "No se pudo leer el archivo .xlsx. Verificar que sea un Excel valido.",
                code="invalid_xlsx",
            ) from exc

        try:
            sheet = workbook.active
            rows = []
            for row in sheet.iter_rows(values_only=True):
                rows.append(list(row))
                if len(rows) > MAX_IMPORT_ROWS + 1:
                    break
            return rows
        finally:
            workbook.close()

    def _read_xls(self, upload):
        try:
            import xlrd
        except ModuleNotFoundError as exc:
            raise ProductImportFormatError(
                "El soporte para .xls requiere instalar xlrd.",
                code="missing_dependency",
            ) from exc

        raw = upload.read()
        if not raw:
            raise ProductImportFormatError("El archivo esta vacio.", code="empty_file")

        try:
            workbook = xlrd.open_workbook(file_contents=raw)
            sheet = workbook.sheet_by_index(0)
        except Exception as exc:
            raise ProductImportFormatError(
                "No se pudo leer el archivo .xls. Verificar que sea un Excel valido.",
                code="invalid_xls",
            ) from exc

        max_rows = min(sheet.nrows, MAX_IMPORT_ROWS + 2)
        return [
            [sheet.cell_value(row_index, col_index) for col_index in range(sheet.ncols)]
            for row_index in range(max_rows)
        ]

    def _table_rows_to_dicts(self, table_rows):
        non_empty_rows = [
            (index, row)
            for index, row in enumerate(table_rows, start=1)
            if not self._is_empty_row(row)
        ]
        if not non_empty_rows:
            raise ProductImportFormatError("El archivo esta vacio.", code="empty_file")

        header_number, header_row = non_empty_rows[0]
        header_map, header_errors = self._build_header_map(header_number, header_row)
        if header_errors:
            raise ProductImportFormatError(
                "El archivo tiene columnas invalidas.",
                code=header_errors[0]["code"],
                errors=header_errors,
            )

        rows = []
        for row_number, row in non_empty_rows[1:]:
            row_data = {"__row__": row_number}
            for column_index, canonical_name in header_map.items():
                row_data[canonical_name] = row[column_index] if column_index < len(row) else ""
            rows.append(row_data)

        if not rows:
            raise ProductImportFormatError(
                "El archivo no contiene filas de productos.",
                code="no_rows",
            )
        if len(rows) > MAX_IMPORT_ROWS:
            raise ProductImportFormatError(
                f"El archivo supera el maximo de {MAX_IMPORT_ROWS} filas.",
                code="too_many_rows",
            )
        return rows

    def _build_header_map(self, header_number, header_row):
        header_map = {}
        mapped_names = {}
        errors = []

        for column_index, raw_header in enumerate(header_row):
            normalized = _normalize_key(raw_header)
            if not normalized:
                continue
            canonical_name = COLUMN_ALIASES.get(normalized)
            if canonical_name is None:
                continue
            if canonical_name in mapped_names:
                errors.append(_row_error(
                    header_number,
                    canonical_name,
                    f"Columna duplicada para '{canonical_name}'.",
                    code="duplicate_column",
                ))
                continue
            header_map[column_index] = canonical_name
            mapped_names[canonical_name] = column_index

        missing_columns = sorted(REQUIRED_COLUMNS - set(mapped_names))
        for column in missing_columns:
            errors.append(_row_error(
                header_number,
                column,
                f"Falta la columna requerida '{column}'.",
                code="missing_columns",
            ))

        return header_map, errors

    def _is_empty_row(self, row):
        return not any(_clean_text(value) for value in row)

    def _normalize_product_data(self, tenant_id, row):
        row_number = row["__row__"]
        errors = []

        name = _clean_text(row.get("name"), 200)
        if not name:
            errors.append(_row_error(row_number, "name", "El nombre es requerido.", "required"))

        price = self._parse_decimal(row.get("price"), "price", row_number, errors, required=True, decimal_places=2)
        cost = self._parse_decimal(row.get("cost"), "cost", row_number, errors, default=Decimal("0.00"), decimal_places=2)
        stock = self._parse_decimal(row.get("stock"), "stock", row_number, errors, default=Decimal("0.000"), decimal_places=3)
        min_stock = self._parse_decimal(row.get("min_stock"), "min_stock", row_number, errors, default=Decimal("0.000"), decimal_places=3)
        unit = self._normalize_unit(row.get("unit"), row_number, errors)
        category_ref = self._normalize_category_ref(row, row_number, errors)

        return {
            "name": name,
            "description": _clean_text(row.get("description")),
            "price": price,
            "cost": cost,
            "unit": unit,
            "stock": stock,
            "min_stock": min_stock,
            "barcode": _clean_text(row.get("barcode"), 100),
            "image_url": _clean_text(row.get("image_url"), 500),
            "_category_ref": category_ref,
        }, errors

    def _parse_decimal(self, value, field, row_number, errors, default=None, required=False, decimal_places=2):
        text = _clean_text(value)
        if not text:
            if required:
                errors.append(_row_error(row_number, field, f"El campo '{field}' es requerido.", "required"))
                return None
            return default

        normalized = text.replace("$", "").replace(" ", "")
        if "," in normalized and "." in normalized:
            if normalized.rfind(",") > normalized.rfind("."):
                normalized = normalized.replace(".", "").replace(",", ".")
            else:
                normalized = normalized.replace(",", "")
        elif "," in normalized:
            normalized = normalized.replace(",", ".")

        try:
            amount = Decimal(normalized)
        except (InvalidOperation, ValueError):
            errors.append(_row_error(row_number, field, f"El valor '{text}' no es un numero valido.", "invalid_decimal"))
            return default

        if not amount.is_finite():
            errors.append(_row_error(row_number, field, f"El valor '{text}' no es un numero finito.", "invalid_decimal"))
            return default

        if amount < 0:
            errors.append(_row_error(row_number, field, "El valor no puede ser negativo.", "negative"))
            return default

        quantizer = Decimal("1").scaleb(-decimal_places)
        try:
            amount = amount.quantize(quantizer)
        except InvalidOperation:
            errors.append(_row_error(row_number, field, f"El valor '{text}' no respeta la precision permitida.", "invalid_decimal"))
            return default

        max_digits, model_decimal_places = PRODUCT_DECIMAL_LIMITS[field]
        if not self._fits_decimal_field(amount, max_digits, model_decimal_places):
            errors.append(_row_error(
                row_number,
                field,
                f"El valor supera el maximo permitido para '{field}'.",
                "max_digits",
            ))
            return default

        return amount

    def _fits_decimal_field(self, amount, max_digits, decimal_places):
        sign, digits, exponent = amount.as_tuple()
        decimals = abs(exponent) if exponent < 0 else 0
        integer_digits = len(digits) - decimals
        if integer_digits < 0:
            integer_digits = 0
        return decimals <= decimal_places and integer_digits <= max_digits - decimal_places

    def _normalize_unit(self, value, row_number, errors):
        text = _normalize_key(value)
        if not text:
            return Product.Unit.U
        if text in {"u", "unit", "unidad", "un", "unidad_es"}:
            return Product.Unit.U
        if text in {"kg", "kilo", "kilos", "kilogram", "kilogramo", "kilogramos"}:
            return Product.Unit.KG
        errors.append(_row_error(
            row_number,
            "unit",
            "Unidad invalida. Valores permitidos: U, KG.",
            "invalid_choice",
        ))
        return Product.Unit.U

    def _normalize_category_ref(self, row, row_number, errors):
        category_id = _clean_text(row.get("category_id"))
        category_name = _clean_text(row.get("category"), 100)

        if category_id:
            try:
                parsed_uuid = uuid.UUID(category_id)
            except ValueError:
                errors.append(_row_error(row_number, "category_id", "category_id debe ser un UUID valido.", "invalid_uuid"))
                return None
            return {"id": parsed_uuid}

        if category_name:
            return {"name": category_name}

        return None

    def _resolve_category(self, tenant_id, category_ref, row_number):
        if not category_ref:
            return None, []

        if category_ref.get("id"):
            parsed_uuid = category_ref["id"]
            category = Category.objects.filter(uuid=parsed_uuid, tenant_id=tenant_id).first()
            if category is None:
                return None, [_row_error(row_number, "category_id", "Categoria no encontrada para este tenant.", "not_found")]
            return category, []

        category_name = category_ref["name"]
        category = Category.all_objects.filter(
            tenant_id=tenant_id,
            name=category_name,
        ).order_by("-active", "-updated_at").first()
        if category:
            if not category.active:
                category.active = True
                category.save(update_fields=["active", "updated_at"])
            return category, []

        try:
            return Category.objects.create(tenant_id=tenant_id, name=category_name), []
        except IntegrityError:
            return Category.objects.get(tenant_id=tenant_id, name=category_name), []

    def _preview_category_errors(self, tenant_id, category_ref, row_number):
        if not category_ref or not category_ref.get("id"):
            return []
        category = Category.objects.filter(uuid=category_ref["id"], tenant_id=tenant_id).first()
        if category is None:
            return [_row_error(row_number, "category_id", "Categoria no encontrada para este tenant.", "not_found")]
        return []

    def _validate_file_duplicates(self, row_number, product_data, seen_file_keys):
        errors = []
        for key, field in self._row_unique_keys(product_data):
            first_row = seen_file_keys.get(key)
            if first_row is not None:
                errors.append(_row_error(
                    row_number,
                    field,
                    f"Duplicado dentro del archivo. Primera aparicion: fila {first_row}.",
                    "duplicate_in_file",
                ))
        return errors

    def _remember_file_keys(self, row_number, product_data, seen_file_keys):
        for key, _field in self._row_unique_keys(product_data):
            seen_file_keys[key] = row_number

    def _row_unique_keys(self, product_data):
        keys = [(f"name:{product_data['name']}", "name")]
        if product_data["barcode"]:
            keys.append((f"barcode:{product_data['barcode']}", "barcode"))
        return keys

    def _match_active_product(self, tenant_id, product_data, row_number):
        queryset = Product.all_objects.filter(tenant_id=tenant_id, active=True)
        barcode_match = None
        if product_data["barcode"]:
            barcode_match = queryset.filter(barcode=product_data["barcode"]).first()
        name_match = queryset.filter(name=product_data["name"]).first()

        if barcode_match and name_match and barcode_match.uuid != name_match.uuid:
            return {
                "product": None,
                "match_type": None,
                "errors": [_row_error(
                    row_number,
                    "barcode",
                    "El codigo de barras coincide con un producto y el nombre con otro.",
                    "match_conflict",
                )],
            }

        product = barcode_match or name_match
        match_type = "barcode" if barcode_match else "name" if name_match else None
        return {
            "product": product,
            "match_type": match_type,
            "errors": [],
        }

    def _upsert_product(self, tenant_id, product_data, row_number):
        match = self._match_active_product(tenant_id, product_data, row_number)
        if match["errors"]:
            return {
                "product": None,
                "action": None,
                "match_type": None,
                "errors": match["errors"],
            }
        product = match["product"]
        match_type = match["match_type"]

        try:
            with transaction.atomic():
                product_fields = product_data.copy()
                category_ref = product_fields.pop("_category_ref", None)
                category, category_errors = self._resolve_category(tenant_id, category_ref, row_number)
                if category_errors:
                    return {
                        "product": None,
                        "action": None,
                        "match_type": None,
                        "errors": category_errors,
                    }
                product_fields["category"] = category

                if product is None:
                    product = Product.all_objects.create(
                        tenant_id=tenant_id,
                        active=True,
                        **product_fields,
                    )
                    action = "created"
                else:
                    for field, value in product_fields.items():
                        setattr(product, field, value)
                    product.active = True
                    product.save()
                    action = "updated"
        except IntegrityError:
            return {
                "product": None,
                "action": None,
                "match_type": None,
                "errors": [_row_error(
                    row_number,
                    None,
                    "No se pudo guardar el producto por una restriccion de unicidad.",
                    "integrity_error",
                )],
            }

        return {
            "product": product,
            "action": action,
            "match_type": match_type,
            "errors": [],
        }
